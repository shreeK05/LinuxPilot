from typing import Any, Dict
import os
from pypdf import PdfReader
from openpyxl import load_workbook, Workbook
from app.agent.actions.registry import ActionHandler, ActionExecutionResult
from app.agent.models import ActionDefinition

class PDFExtractTextHandler(ActionHandler):
    """
    Extracts text from a PDF file.
    Parameters:
      - path: str
    """
    def execute(self, action: ActionDefinition) -> ActionExecutionResult:
        try:
            path = action.parameters.get("path")
            if not path or not os.path.exists(path):
                return ActionExecutionResult(success=False, output=None, error=f"File not found: {path}")

            reader = PdfReader(path)
            text_content = ""
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_content += extracted + "\n"
                    
            return ActionExecutionResult(success=True, output={"text": text_content.strip()})
        except Exception as e:
            return ActionExecutionResult(success=False, output=None, error=f"Failed to extract PDF text: {str(e)}")

class XLSXReadHandler(ActionHandler):
    """
    Reads rows from an XLSX file.
    Parameters:
      - path: str
      - sheet_name: str (optional)
    """
    def execute(self, action: ActionDefinition) -> ActionExecutionResult:
        try:
            path = action.parameters.get("path")
            if not path or not os.path.exists(path):
                return ActionExecutionResult(success=False, output=None, error=f"File not found: {path}")

            sheet_name = action.parameters.get("sheet_name")
            
            wb = load_workbook(path, data_only=True)
            if sheet_name and sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
            else:
                sheet = wb.active
                
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(list(row))
                
            return ActionExecutionResult(success=True, output={"rows": data})
        except Exception as e:
            return ActionExecutionResult(success=False, output=None, error=f"Failed to read XLSX: {str(e)}")

class XLSXWriteHandler(ActionHandler):
    """
    Writes rows to an XLSX file.
    Parameters:
      - path: str
      - rows: List[List[Any]]
      - sheet_name: str (optional)
    """
    def execute(self, action: ActionDefinition) -> ActionExecutionResult:
        try:
            path = action.parameters.get("path")
            rows = action.parameters.get("rows", [])
            sheet_name = action.parameters.get("sheet_name", "Sheet1")
            
            if not path:
                return ActionExecutionResult(success=False, output=None, error="Path is required")

            if os.path.exists(path):
                wb = load_workbook(path)
                if sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                else:
                    sheet = wb.create_sheet(title=sheet_name)
            else:
                wb = Workbook()
                sheet = wb.active
                sheet.title = sheet_name

            for row_data in rows:
                sheet.append(row_data)

            wb.save(path)
                
            return ActionExecutionResult(success=True, output={"path": path, "rows_written": len(rows)})
        except Exception as e:
            return ActionExecutionResult(success=False, output=None, error=f"Failed to write XLSX: {str(e)}")
